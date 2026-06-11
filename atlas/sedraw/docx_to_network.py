"""Convert a Word (.docx) network/TD intake into the SEdraw network-topology structure with the local
Nemotron model — a best-effort map onto the Excel intake schema the network generator already consumes.

Pipeline (driven by service.convert_docx_network):
    .docx → extract_text() → parse_to_network() (Nemotron) → {"sites":[…]} → _clean_network →
    write_network_topology (.xlsx) → generate_network (.drawio + viewer)

Not everything in a Word doc maps cleanly; the model maps what it can and omits the rest. The user can
then edit the generated intake .xlsx in the GUI editor and regenerate.
"""
from __future__ import annotations

from pathlib import Path

from atlas.engine import nemotron
from atlas.engine.extract import extract_json
from atlas.sedraw import constants as C

_CAP = 6000   # chars of doc text fed to the model (leaves room for Nemotron's reasoning + the JSON output)


def extract_text(path) -> str:
    """Flatten a .docx into plain text: paragraphs + every table (rows as 'cell | cell'). Capped."""
    import docx
    doc = docx.Document(str(path))
    lines: list[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            lines.append(t)
    for ti, tb in enumerate(doc.tables):
        lines.append(f"\n[TABLE {ti + 1}]")
        for row in tb.rows:
            cells = [(c.text or "").strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    text = "\n".join(lines).strip()
    if len(text) > _CAP:
        text = text[:_CAP] + "\n…(document truncated to fit the model context)…"
    return text


def _prompt(text: str) -> str:
    cats = ", ".join(C.NET_CATEGORIES)
    return (
        "You convert a customer's network / technical-discovery intake (extracted from a Word document) "
        "into a STRUCTURED network topology. Map what you reasonably can and OMIT anything that doesn't "
        "fit — do not invent devices.\n\n"
        "Return ONLY this JSON (no prose, no code fences):\n"
        '{"sites":[{"label":"<site name>","site_type":"dc"|"branch",'
        '"categories":[{"category":"<one of the allowed categories>","qty_model":"<qty x model>",'
        '"connected_to":"<what it uplinks/connects to>","notes":"<short note or empty>"}]}]}\n\n'
        f"ALLOWED categories (use the closest match; skip anything that fits none): {cats}\n\n"
        "Rules:\n"
        "- One entry in \"sites\" per physical site (Primary/Core/DC → site_type \"dc\"; "
        "Secondary/DR is also \"dc\"; Remote/Branch → \"branch\").\n"
        "- Put quantities and model names in qty_model (e.g. \"2x Palo Alto PA-850 (HA)\").\n"
        "- Put uplinks / what-connects-to-what in connected_to.\n"
        "- Group branch locations that share the same infrastructure pattern into one branch site if the "
        "doc says so.\n"
        "- If you can't determine a topology, return {\"sites\":[]}.\n\n"
        f"=== INTAKE DOCUMENT ===\n{text}\n\n=== END ===\nJSON:")


def _coerce(d) -> dict:
    """Coerce the model output into a {'sites':[...]} dict the network writer accepts."""
    if not isinstance(d, dict):
        return {"sites": []}
    sites = d.get("sites")
    if not isinstance(sites, list):
        return {"sites": []}
    valid = {c.lower(): c for c in C.NET_CATEGORIES}
    out_sites = []
    for s in sites:
        if not isinstance(s, dict):
            continue
        cats = []
        for c in (s.get("categories") or []):
            if not isinstance(c, dict):
                continue
            name = (c.get("category") or "").strip()
            canon = valid.get(name.lower(), name)        # snap to a known category name when possible
            if not canon:
                continue
            cats.append({
                "category": canon,
                "qty_model": (c.get("qty_model") or "").strip(),
                "connected_to": (c.get("connected_to") or "").strip(),
                "notes": (c.get("notes") or "").strip(),
            })
        st = (s.get("site_type") or "branch").strip().lower()
        out_sites.append({"label": (s.get("label") or "Site").strip(),
                          "site_type": "dc" if st == "dc" else "branch",
                          "categories": cats})
    return {"sites": out_sites}


def parse_to_network(text: str, on_log=None, should_cancel=None) -> dict:
    """Nemotron maps the intake text → {'sites':[…]}. Fail-soft to {'sites':[]}."""
    if not (text or "").strip():
        return {"sites": []}
    if on_log:
        on_log("🧠 Nemotron is reading the Word document and mapping it to a network intake…")
    try:
        # Generous output budget: Nemotron emits a long <think> block before the JSON; too small a cap
        # truncates the JSON mid-object. Input is capped (_CAP) so input + this fits the 8K window.
        raw = nemotron.complete(_prompt(text), on_log=on_log, max_tokens=6000, should_cancel=should_cancel)
        return _coerce(extract_json(raw))
    except nemotron.Cancelled:
        raise
    except Exception as e:  # noqa: BLE001
        if on_log:
            on_log(f"⚠ Could not parse the document ({str(e)[:120]}).")
        return {"sites": []}
